#!/usr/bin/env python3
import csv,io,re,urllib.request
from pathlib import Path
from openpyxl import load_workbook
PAGE='https://www.jpx.co.jp/markets/statistics-equities/misc/01.html'
UA='Mozilla/5.0'
html=urllib.request.urlopen(urllib.request.Request(PAGE,headers={'User-Agent':UA}),timeout=30).read().decode('utf-8','ignore')
links=re.findall(r'href="([^"]+\.xlsx[^"]*)"',html,re.I)
if not links: raise SystemExit('JPX xlsx link not found')
url=links[0]
if url.startswith('/'): url='https://www.jpx.co.jp'+url
raw=urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':UA}),timeout=60).read()
wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); ws=wb.active
rows=list(ws.iter_rows(values_only=True)); hdr=[str(x or '').strip() for x in rows[0]]
def col(keys):
 for i,h in enumerate(hdr):
  if any(k in h for k in keys): return i
 raise RuntimeError(str(keys))
ic=col(['コード']); inn=col(['銘柄名']); im=col(['市場・商品区分','市場区分'])
out=[]
for r in rows[1:]:
 code=str(r[ic] or '').strip().replace('.0',''); name=str(r[inn] or '').strip(); market=str(r[im] or '')
 if code and name and 'プライム' in market and '外国' not in market:
  out.append({'code':code,'company_name':name,'ticker_tse':f'{code}.T'})
if not (1400 <= len(out) <= 1700): raise SystemExit(f'unexpected Prime count={len(out)}')
p=Path('data/constituents/tse_prime/current.csv'); p.parent.mkdir(parents=True,exist_ok=True)
with p.open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=['code','company_name','ticker_tse']); w.writeheader(); w.writerows(out)
print(f'JPX Prime universe COMPLETE count={len(out)} source={url}')
